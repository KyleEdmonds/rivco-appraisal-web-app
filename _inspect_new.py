import openpyxl, sys

def dump(path, maxrows=8):
    print("="*80)
    print("FILE:", path)
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        print("SHEET:", repr(ws.title), "rows", ws.max_row, "cols", ws.max_column)
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        hdr = rows[0]
        print("HEADERS (%d):" % len(hdr))
        for i, h in enumerate(hdr):
            print("  [%d] %r" % (i, h))
        print("--- first %d data rows ---" % maxrows)
        for r in rows[1:1+maxrows]:
            print("ROW:")
            for i, v in enumerate(r):
                if i < len(hdr) and v is not None:
                    print("   %s = %r" % (hdr[i], v))

for p in sys.argv[1:]:
    dump(p)
